"""분기 사업분석 탭 — 엑셀 레이아웃·수식은 그대로 두고, 라벨별 대시보드 값만 반영한다."""
from __future__ import annotations

import html
import os
import re
import shutil
from datetime import date
from typing import Any

import pandas as pd
import streamlit as st

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

QA_DIR = os.path.join("uploaded_cache", "quarterly_analysis")
QA_TEMPLATE = os.path.join(QA_DIR, "template.xlsx")
QA_WORK = os.path.join(QA_DIR, "work.xlsx")
QA_SHEET = "초안_v2"

# 엑셀에 쓰는 입력 칸만. 수식 칸(K6/K7/K9, D26…)·서술 칸은 건드리지 않는다.
# month = 분기 안 슬롯(1·2·3). 실제 달력 월은 선택한 분기로 매핑한다.
# kind: goal | sales_mil | prior_sales_mil | rate_vs_goal | rate_vs_prior | qty_ton | qty_bt | ...
_GOAL_CELLS = ("E6", "G6", "I6")
_QA_FIELDS: list[dict[str, Any]] = [
    # ■ 목표 / 달성 — 목표는 사용자가 직접 입력
    {"cell": "E6", "label": "목표 · {m1}", "kind": "goal", "month": 1},
    {"cell": "G6", "label": "목표 · {m2}", "kind": "goal", "month": 2},
    {"cell": "I6", "label": "목표 · {m3}", "kind": "goal", "month": 3},
    {"cell": "E7", "label": "실적 · {m1}", "kind": "sales_mil", "month": 1},
    {"cell": "G7", "label": "실적 · {m2}", "kind": "sales_mil", "month": 2},
    {"cell": "I7", "label": "실적 · {m3}", "kind": "sales_mil", "month": 3},
    {"cell": "E8", "label": "달성률 · {m1}", "kind": "rate_vs_goal", "month": 1, "goal_cell": "E6"},
    {"cell": "G8", "label": "달성률 · {m2}", "kind": "rate_vs_goal", "month": 2, "goal_cell": "G6"},
    {"cell": "I8", "label": "달성률 · {m3}", "kind": "rate_vs_goal", "month": 3, "goal_cell": "I6"},
    {"cell": "K8", "label": "달성률 · 합계", "kind": "rate_vs_goal_sum"},
    {"cell": "E9", "label": "전년 매출 · {m1}", "kind": "prior_sales_mil", "month": 1},
    {"cell": "G9", "label": "전년 매출 · {m2}", "kind": "prior_sales_mil", "month": 2},
    {"cell": "I9", "label": "전년 매출 · {m3}", "kind": "prior_sales_mil", "month": 3},
    {"cell": "E10", "label": "전년 비교 · {m1}", "kind": "rate_vs_prior", "month": 1},
    {"cell": "G10", "label": "전년 비교 · {m2}", "kind": "rate_vs_prior", "month": 2},
    {"cell": "I10", "label": "전년 비교 · {m3}", "kind": "rate_vs_prior", "month": 3},
    {"cell": "K10", "label": "전년 비교 · 합계", "kind": "rate_vs_prior_sum"},
    # ■ 실적 / 분석 — 당해·전년 수량 (TON / BT). Bulk 합계·전년대비 일부는 수식/텍스트 유지
    {"cell": "D21", "label": "L-O2 · 당해 {m1}", "kind": "qty_ton", "month": 1, "item": "L-O2", "year": "curr"},
    {"cell": "E21", "label": "L-O2 · 전년 {m1}", "kind": "qty_ton", "month": 1, "item": "L-O2", "year": "prior"},
    {"cell": "F21", "label": "L-O2 · 전년대비 {m1}", "kind": "qty_rate", "month": 1, "item": "L-O2"},
    {"cell": "G21", "label": "L-O2 · 당해 {m2}", "kind": "qty_ton", "month": 2, "item": "L-O2", "year": "curr"},
    {"cell": "H21", "label": "L-O2 · 전년 {m2}", "kind": "qty_ton", "month": 2, "item": "L-O2", "year": "prior"},
    {"cell": "I21", "label": "L-O2 · 전년대비 {m2}", "kind": "qty_rate", "month": 2, "item": "L-O2"},
    {"cell": "J21", "label": "L-O2 · 당해 {m3}", "kind": "qty_ton", "month": 3, "item": "L-O2", "year": "curr"},
    {"cell": "K21", "label": "L-O2 · 전년 {m3}", "kind": "qty_ton", "month": 3, "item": "L-O2", "year": "prior"},
    {"cell": "L21", "label": "L-O2 · 전년대비 {m3}", "kind": "qty_rate", "month": 3, "item": "L-O2"},
    {"cell": "D22", "label": "L-N2 · 당해 {m1}", "kind": "qty_ton", "month": 1, "item": "L-N2", "year": "curr"},
    {"cell": "E22", "label": "L-N2 · 전년 {m1}", "kind": "qty_ton", "month": 1, "item": "L-N2", "year": "prior"},
    {"cell": "F22", "label": "L-N2 · 전년대비 {m1}", "kind": "qty_rate", "month": 1, "item": "L-N2"},
    {"cell": "G22", "label": "L-N2 · 당해 {m2}", "kind": "qty_ton", "month": 2, "item": "L-N2", "year": "curr"},
    {"cell": "H22", "label": "L-N2 · 전년 {m2}", "kind": "qty_ton", "month": 2, "item": "L-N2", "year": "prior"},
    {"cell": "I22", "label": "L-N2 · 전년대비 {m2}", "kind": "qty_rate", "month": 2, "item": "L-N2"},
    {"cell": "J22", "label": "L-N2 · 당해 {m3}", "kind": "qty_ton", "month": 3, "item": "L-N2", "year": "curr"},
    {"cell": "K22", "label": "L-N2 · 전년 {m3}", "kind": "qty_ton", "month": 3, "item": "L-N2", "year": "prior"},
    {"cell": "L22", "label": "L-N2 · 전년대비 {m3}", "kind": "qty_rate", "month": 3, "item": "L-N2"},
    {"cell": "D23", "label": "L-AR · 당해 {m1}", "kind": "qty_ton", "month": 1, "item": "L-AR", "year": "curr"},
    {"cell": "E23", "label": "L-AR · 전년 {m1}", "kind": "qty_ton", "month": 1, "item": "L-AR", "year": "prior"},
    {"cell": "F23", "label": "L-AR · 전년대비 {m1}", "kind": "qty_rate", "month": 1, "item": "L-AR"},
    {"cell": "G23", "label": "L-AR · 당해 {m2}", "kind": "qty_ton", "month": 2, "item": "L-AR", "year": "curr"},
    {"cell": "H23", "label": "L-AR · 전년 {m2}", "kind": "qty_ton", "month": 2, "item": "L-AR", "year": "prior"},
    {"cell": "I23", "label": "L-AR · 전년대비 {m2}", "kind": "qty_rate", "month": 2, "item": "L-AR"},
    {"cell": "J23", "label": "L-AR · 당해 {m3}", "kind": "qty_ton", "month": 3, "item": "L-AR", "year": "curr"},
    {"cell": "K23", "label": "L-AR · 전년 {m3}", "kind": "qty_ton", "month": 3, "item": "L-AR", "year": "prior"},
    {"cell": "L23", "label": "L-AR · 전년대비 {m3}", "kind": "qty_rate", "month": 3, "item": "L-AR"},
    {"cell": "D24", "label": "L-CO2 · 당해 {m1}", "kind": "qty_ton", "month": 1, "item": "L-CO2", "year": "curr"},
    {"cell": "E24", "label": "L-CO2 · 전년 {m1}", "kind": "qty_ton", "month": 1, "item": "L-CO2", "year": "prior"},
    {"cell": "F24", "label": "L-CO2 · 전년대비 {m1}", "kind": "qty_rate", "month": 1, "item": "L-CO2"},
    {"cell": "G24", "label": "L-CO2 · 당해 {m2}", "kind": "qty_ton", "month": 2, "item": "L-CO2", "year": "curr"},
    {"cell": "H24", "label": "L-CO2 · 전년 {m2}", "kind": "qty_ton", "month": 2, "item": "L-CO2", "year": "prior"},
    {"cell": "I24", "label": "L-CO2 · 전년대비 {m2}", "kind": "qty_rate", "month": 2, "item": "L-CO2"},
    {"cell": "J24", "label": "L-CO2 · 당해 {m3}", "kind": "qty_ton", "month": 3, "item": "L-CO2", "year": "curr"},
    {"cell": "K24", "label": "L-CO2 · 전년 {m3}", "kind": "qty_ton", "month": 3, "item": "L-CO2", "year": "prior"},
    {"cell": "L24", "label": "L-CO2 · 전년대비 {m3}", "kind": "qty_rate", "month": 3, "item": "L-CO2"},
    {"cell": "D25", "label": "LPG · 당해 {m1}", "kind": "qty_ton", "month": 1, "item": "LPG", "year": "curr"},
    {"cell": "E25", "label": "LPG · 전년 {m1}", "kind": "qty_ton", "month": 1, "item": "LPG", "year": "prior"},
    {"cell": "F25", "label": "LPG · 전년대비 {m1}", "kind": "qty_rate", "month": 1, "item": "LPG"},
    {"cell": "G25", "label": "LPG · 당해 {m2}", "kind": "qty_ton", "month": 2, "item": "LPG", "year": "curr"},
    {"cell": "H25", "label": "LPG · 전년 {m2}", "kind": "qty_ton", "month": 2, "item": "LPG", "year": "prior"},
    {"cell": "I25", "label": "LPG · 전년대비 {m2}", "kind": "qty_rate", "month": 2, "item": "LPG"},
    {"cell": "J25", "label": "LPG · 당해 {m3}", "kind": "qty_ton", "month": 3, "item": "LPG", "year": "curr"},
    {"cell": "K25", "label": "LPG · 전년 {m3}", "kind": "qty_ton", "month": 3, "item": "LPG", "year": "prior"},
    {"cell": "L25", "label": "LPG · 전년대비 {m3}", "kind": "qty_rate", "month": 3, "item": "LPG"},
    {"cell": "F26", "label": "Bulk 합계 · 전년대비 {m1}", "kind": "bulk_rate", "month": 1},
    {"cell": "I26", "label": "Bulk 합계 · 전년대비 {m2}", "kind": "bulk_rate", "month": 2},
    {"cell": "L26", "label": "Bulk 합계 · 전년대비 {m3}", "kind": "bulk_rate", "month": 3},
    {"cell": "D27", "label": "Gas Cylinder · 당해 {m1}", "kind": "qty_bt", "month": 1, "year": "curr"},
    {"cell": "E27", "label": "Gas Cylinder · 전년 {m1}", "kind": "qty_bt", "month": 1, "year": "prior"},
    {"cell": "F27", "label": "Gas Cylinder · 전년대비 {m1}", "kind": "cyl_rate", "month": 1},
    {"cell": "G27", "label": "Gas Cylinder · 당해 {m2}", "kind": "qty_bt", "month": 2, "year": "curr"},
    {"cell": "H27", "label": "Gas Cylinder · 전년 {m2}", "kind": "qty_bt", "month": 2, "year": "prior"},
    {"cell": "I27", "label": "Gas Cylinder · 전년대비 {m2}", "kind": "cyl_rate", "month": 2},
    {"cell": "J27", "label": "Gas Cylinder · 당해 {m3}", "kind": "qty_bt", "month": 3, "year": "curr"},
    {"cell": "K27", "label": "Gas Cylinder · 전년 {m3}", "kind": "qty_bt", "month": 3, "year": "prior"},
    {"cell": "L27", "label": "Gas Cylinder · 전년대비 {m3}", "kind": "cyl_rate", "month": 3},
]

_BULK_ITEMS = ("L-O2", "L-N2", "L-AR", "L-CO2", "LPG")
_QA_QUARTER_MONTHS = {
    1: (1, 2, 3),
    2: (4, 5, 6),
    3: (7, 8, 9),
    4: (10, 11, 12),
}


def _qa_quarter_options(*, start_year: int = 2020, end_year: int | None = None) -> list[str]:
    """연도·분기 목록 — selectbox에서 스크롤로 고른다."""
    today = date.today()
    end = int(end_year or (today.year + 1))
    opts: list[str] = []
    for y in range(int(start_year), end + 1):
        for q in (1, 2, 3, 4):
            opts.append(f"{y}년 {q}분기")
    return opts


def _qa_parse_quarter_label(label: str) -> tuple[int, int]:
    raw = str(label or "").strip()
    m = re.match(r"(\d{4})\s*년\s*([1-4])\s*분기", raw)
    if not m:
        today = date.today()
        q = (today.month - 1) // 3 + 1
        return today.year, q
    return int(m.group(1)), int(m.group(2))


def _qa_default_quarter_label() -> str:
    today = date.today()
    q = (today.month - 1) // 3 + 1
    # 분기 초·중이면 직전 분기 기본
    if today.day < 10 and today.month in (1, 4, 7, 10):
        if q == 1:
            return f"{today.year - 1}년 4분기"
        return f"{today.year}년 {q - 1}분기"
    return f"{today.year}년 {q}분기"


def _qa_slot_month(months: tuple[int, int, int], slot: int) -> int:
    idx = max(1, min(3, int(slot))) - 1
    return int(months[idx])


def _qa_field_label(tmpl: str, months: tuple[int, int, int]) -> str:
    return (
        str(tmpl)
        .replace("{m1}", f"{months[0]}월")
        .replace("{m2}", f"{months[1]}월")
        .replace("{m3}", f"{months[2]}월")
    )


def _qa_ensure_workfile() -> str:
    os.makedirs(QA_DIR, exist_ok=True)
    if not os.path.isfile(QA_WORK):
        src = QA_TEMPLATE if os.path.isfile(QA_TEMPLATE) else None
        if not src:
            raise FileNotFoundError(
                f"분기 사업분석 템플릿이 없습니다: {QA_TEMPLATE}"
            )
        shutil.copy2(src, QA_WORK)
    return QA_WORK


def _qa_read_cells(path: str, cells: list[str]) -> dict[str, Any]:
    if load_workbook is None:
        raise RuntimeError("openpyxl 이 필요합니다.")
    wb = load_workbook(path, data_only=False)
    if QA_SHEET not in wb.sheetnames:
        wb.close()
        raise KeyError(f"시트 '{QA_SHEET}' 없음: {wb.sheetnames}")
    ws = wb[QA_SHEET]
    out = {addr: ws[addr].value for addr in cells}
    wb.close()
    return out


def _qa_write_cells(path: str, values: dict[str, Any]) -> int:
    """입력 칸만 덮어쓴다. 수식·병합·서식은 유지."""
    if load_workbook is None:
        raise RuntimeError("openpyxl 이 필요합니다.")
    wb = load_workbook(path, data_only=False)
    ws = wb[QA_SHEET]
    n = 0
    for addr, val in values.items():
        cell = ws[addr]
        cur = cell.value
        if isinstance(cur, str) and cur.startswith("="):
            continue
        if cur == val:
            continue
        cell.value = val
        n += 1
    wb.save(path)
    wb.close()
    return n


def _qa_arrow_pct(curr: float, base: float) -> str:
    if base == 0:
        return "-" if curr == 0 else "▲"
    pct = round((curr - base) / abs(base) * 100)
    if pct > 0:
        return f"▲{pct}%"
    if pct < 0:
        return f"▼{abs(pct)}%"
    return "0%"


def _qa_prepare_df(df: pd.DataFrame | None) -> pd.DataFrame:
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return pd.DataFrame()
    work = df.copy()
    if "매출일_dt" not in work.columns and "매출일" in work.columns:
        work["매출일_dt"] = pd.to_datetime(work["매출일"], errors="coerce")
    if "매출일_dt" not in work.columns:
        return pd.DataFrame()
    work["매출일_dt"] = pd.to_datetime(work["매출일_dt"], errors="coerce")
    if "매출액" in work.columns:
        work["매출액"] = pd.to_numeric(work["매출액"], errors="coerce").fillna(0.0)
    else:
        work["매출액"] = 0.0
    if "출고량" in work.columns:
        work["출고량"] = pd.to_numeric(work["출고량"], errors="coerce").fillna(0.0)
    else:
        work["출고량"] = 0.0
    if "품목명" not in work.columns:
        work["품목명"] = ""
    work = work.dropna(subset=["매출일_dt"])
    work["연도"] = work["매출일_dt"].dt.year
    work["월"] = work["매출일_dt"].dt.month
    return work


def _qa_item_mask(series: pd.Series, item: str) -> pd.Series:
    s = series.astype(str)
    u = s.str.upper().str.replace(" ", "", regex=False)
    is_bulk = u.str.contains("BULK", na=False) | s.str.contains("벌크", na=False)
    if item == "L-O2":
        return is_bulk & ~u.str.contains("CO2", na=False) & (
            u.str.contains("O2", na=False) | s.str.contains("산소", na=False)
        )
    if item == "L-N2":
        return is_bulk & (u.str.contains("N2", na=False) | s.str.contains("질소", na=False))
    if item == "L-AR":
        return is_bulk & (
            u.str.contains("AR", na=False)
            | s.str.contains("아르곤|아르", na=False)
        )
    if item == "L-CO2":
        return is_bulk & (u.str.contains("CO2", na=False) | s.str.contains("탄산", na=False))
    if item == "LPG":
        return is_bulk & (u.str.contains("LPG", na=False) | s.str.contains("엘피지|프로판", na=False))
    return pd.Series(False, index=series.index)


def _qa_is_listed_bulk(series: pd.Series) -> pd.Series:
    """실적/분석 표의 L-O2·L-N2·L-AR·L-CO2·LPG 해당 여부."""
    m = pd.Series(False, index=series.index)
    for item in _BULK_ITEMS:
        m = m | _qa_item_mask(series, item)
    return m


def _qa_is_cylinder(series: pd.Series) -> pd.Series:
    """실적/분석에서 위 5종(L-O2~LPG) 외는 전부 Gas Cylinder(실린더 매출)로 집계."""
    s = series.astype(str).str.strip()
    valid = s.str.len().gt(0) & ~s.isin({"", "nan", "None", "NaN"})
    return valid & ~_qa_is_listed_bulk(series)


def _qa_month_sales_mil(work: pd.DataFrame, year: int, month: int) -> float:
    if work.empty:
        return 0.0
    hit = work[(work["연도"] == year) & (work["월"] == month)]
    # 부가세 별도 백만원 — 매출액(공급가) / 1,000,000
    return float(round(hit["매출액"].sum() / 1_000_000.0))


def _qa_month_qty_ton(work: pd.DataFrame, year: int, month: int, item: str) -> float:
    if work.empty:
        return 0.0
    m = (work["연도"] == year) & (work["월"] == month) & _qa_item_mask(work["품목명"], item)
    # kg → TON
    return float(round(work.loc[m, "출고량"].sum() / 1000.0))


def _qa_month_qty_bt(work: pd.DataFrame, year: int, month: int) -> float:
    if work.empty:
        return 0.0
    m = (work["연도"] == year) & (work["월"] == month) & _qa_is_cylinder(work["품목명"])
    return float(round(work.loc[m, "출고량"].sum()))


def compute_dashboard_values(
    df: pd.DataFrame | None,
    *,
    year: int,
    months: tuple[int, int, int],
    goals: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """라벨(셀)별 대시보드 계산값. 목표는 사용자가 넣은 goals를 쓴다."""
    work = _qa_prepare_df(df)
    prior = year - 1
    goal_map = goals or {}
    m1, m2, m3 = (int(months[0]), int(months[1]), int(months[2]))
    cal_months = (m1, m2, m3)
    sales = {m: _qa_month_sales_mil(work, year, m) for m in cal_months}
    prior_sales = {m: _qa_month_sales_mil(work, prior, m) for m in cal_months}
    ton_curr = {
        (item, m): _qa_month_qty_ton(work, year, m, item)
        for item in _BULK_ITEMS
        for m in cal_months
    }
    ton_prior = {
        (item, m): _qa_month_qty_ton(work, prior, m, item)
        for item in _BULK_ITEMS
        for m in cal_months
    }
    bt_curr = {m: _qa_month_qty_bt(work, year, m) for m in cal_months}
    bt_prior = {m: _qa_month_qty_bt(work, prior, m) for m in cal_months}

    def _goal(cell: str) -> float:
        raw = goal_map.get(cell)
        try:
            return float(raw)
        except (TypeError, ValueError):
            return 0.0

    out: dict[str, Any] = {}
    for f in _QA_FIELDS:
        addr = f["cell"]
        kind = f["kind"]
        slot = int(f.get("month") or 0)
        month = _qa_slot_month(cal_months, slot) if slot else 0
        if kind == "goal":
            out[addr] = int(round(_goal(addr)))
            continue
        if kind == "sales_mil":
            out[addr] = int(sales[month])
            continue
        if kind == "prior_sales_mil":
            out[addr] = int(prior_sales[month])
            continue
        if kind == "rate_vs_goal":
            g = _goal(str(f.get("goal_cell") or ""))
            out[addr] = _qa_arrow_pct(sales[month], g)
            continue
        if kind == "rate_vs_goal_sum":
            gsum = _goal("E6") + _goal("G6") + _goal("I6")
            out[addr] = _qa_arrow_pct(sum(sales.values()), gsum)
            continue
        if kind == "rate_vs_prior":
            out[addr] = _qa_arrow_pct(sales[month], prior_sales[month])
            continue
        if kind == "rate_vs_prior_sum":
            out[addr] = _qa_arrow_pct(sum(sales.values()), sum(prior_sales.values()))
            continue
        item = str(f.get("item") or "")
        if kind == "qty_ton":
            src = ton_curr if f.get("year") == "curr" else ton_prior
            out[addr] = int(src[(item, month)])
            continue
        if kind == "qty_rate":
            out[addr] = _qa_arrow_pct(ton_curr[(item, month)], ton_prior[(item, month)])
            continue
        if kind == "bulk_rate":
            c = sum(ton_curr[(it, month)] for it in _BULK_ITEMS)
            p = sum(ton_prior[(it, month)] for it in _BULK_ITEMS)
            out[addr] = _qa_arrow_pct(c, p)
            continue
        if kind == "qty_bt":
            src = bt_curr if f.get("year") == "curr" else bt_prior
            out[addr] = int(src[month])
            continue
        if kind == "cyl_rate":
            out[addr] = _qa_arrow_pct(bt_curr[month], bt_prior[month])
            continue
    return out


def _qa_num_or(raw: Any, default: float = 0.0) -> float:
    try:
        if raw is None or raw == "":
            return float(default)
        return float(raw)
    except (TypeError, ValueError):
        return float(default)


def _qa_fmt_int(v: Any) -> str:
    try:
        return f"{int(round(float(v))):,}"
    except (TypeError, ValueError):
        return "-"


def _qa_fmt_arrow(v: Any) -> str:
    s = str(v or "-")
    if s.startswith("▲"):
        return f'<span style="color:#d32f2f;font-weight:700">{html.escape(s)}</span>'
    if s.startswith("▼"):
        return f'<span style="color:#1565c0;font-weight:700">{html.escape(s)}</span>'
    return html.escape(s)


def _qa_css() -> str:
    return """
<style>
.qa-wrap { font-family: "Malgun Gothic","Apple SD Gothic Neo",sans-serif; margin: 0 0 1.2rem; }
.qa-head { display:flex; justify-content:space-between; align-items:flex-end; margin: 0 0 .45rem; }
.qa-title { font-size: 1.15rem; font-weight: 800; color:#222; }
.qa-unit { font-size: .85rem; color:#555; }
.qa-table { width:100%; border-collapse:collapse; table-layout:fixed; font-size:14px; }
.qa-table th, .qa-table td { border:1px solid #333; padding:6px 8px; text-align:center; }
.qa-table th { background:#5a5a5a; color:#fff; font-weight:700; }
.qa-table .qa-left { text-align:left; font-weight:700; background:#fafafa; }
.qa-table .qa-year { background:#eee; font-weight:800; vertical-align:middle; }
.qa-table .qa-gray { background:#ececec; }
.qa-table .qa-bulk { background:#fff59d; color:#c62828; font-weight:800; }
.qa-table .qa-bulk td { background:#fff59d; color:#c62828; font-weight:800; }
.qa-qbar { display:flex; gap:.4rem; flex-wrap:wrap; margin:.2rem 0 .8rem; }
</style>
"""


def _qa_goal_table_html(
    year: int,
    months: tuple[int, int, int],
    goals: tuple[int, int, int],
    sales: tuple[int, int, int],
    prior: tuple[int, int, int],
    rates_goal: tuple[str, str, str, str],
    rates_prior: tuple[str, str, str, str],
) -> str:
    m1, m2, m3 = months
    gsum = sum(goals)
    ssum = sum(sales)
    psum = sum(prior)
    rows = [
        ("목표", goals + (gsum,), "", False),
        ("실적", sales + (ssum,), "", False),
        ("달성률", rates_goal[:3] + (rates_goal[3],), "arrow", False),
        ("전년 매출", prior + (psum,), "", True),
        ("전년 비교", rates_prior[:3] + (rates_prior[3],), "arrow", True),
    ]
    body = []
    for i, (lab, vals, kind, gray) in enumerate(rows):
        cls = "qa-gray" if gray else ""
        tds = []
        if i == 0:
            tds.append(
                f'<td class="qa-year" rowspan="3">{html.escape(str(year))}년</td>'
            )
        if i < 3:
            tds.append(f'<td class="qa-left">{html.escape(lab)}</td>')
        else:
            tds.append(f'<td class="qa-left {cls}" colspan="2">{html.escape(lab)}</td>')
        for v in vals:
            if kind == "arrow":
                cell = _qa_fmt_arrow(v)
            else:
                cell = html.escape(_qa_fmt_int(v))
            tds.append(f'<td class="{cls}">{cell}</td>')
        body.append("<tr>" + "".join(tds) + "</tr>")
    return f"""
<div class="qa-wrap">
  <div class="qa-head">
    <div class="qa-title">■ 목표 / 달성</div>
    <div class="qa-unit">(단위 : 백만원, 부가세 별도)</div>
  </div>
  <table class="qa-table">
    <thead>
      <tr>
        <th colspan="2">구 분</th>
        <th>{m1}월</th><th>{m2}월</th><th>{m3}월</th><th>합계</th>
      </tr>
    </thead>
    <tbody>{"".join(body)}</tbody>
  </table>
</div>
"""


def _qa_perf_table_html(
    year: int,
    months: tuple[int, int, int],
    rows: list[dict[str, Any]],
) -> str:
    """rows: label, curr[3], prior[3], rates[3], bulk?"""
    prior_y = year - 1
    y2 = str(year)[2:]
    yp = str(prior_y)[2:]
    m1, m2, m3 = months
    head = f"""
      <tr>
        <th rowspan="2">구분</th>
        <th colspan="3">{m1}월</th>
        <th colspan="3">{m2}월</th>
        <th colspan="3">{m3}월</th>
      </tr>
      <tr>
        <th>'{y2}년 {m1}월</th><th>'{yp}년 {m1}월</th><th>전년 대비</th>
        <th>'{y2}년 {m2}월</th><th>'{yp}년 {m2}월</th><th>전년 대비</th>
        <th>'{y2}년 {m3}월</th><th>'{yp}년 {m3}월</th><th>전년 대비</th>
      </tr>
    """
    body = []
    for r in rows:
        bulk = bool(r.get("bulk"))
        tr_cls = ' class="qa-bulk"' if bulk else ""
        lab = html.escape(str(r.get("label") or ""))
        cells = [f'<td class="qa-left">{lab}</td>']
        curr = list(r.get("curr") or [0, 0, 0])
        prior = list(r.get("prior") or [0, 0, 0])
        rates = list(r.get("rates") or ["-", "-", "-"])
        for i in range(3):
            cells.append(f"<td>{html.escape(_qa_fmt_int(curr[i]))}</td>")
            cells.append(f"<td>{html.escape(_qa_fmt_int(prior[i]))}</td>")
            cells.append(f"<td>{_qa_fmt_arrow(rates[i])}</td>")
        body.append(f"<tr{tr_cls}>" + "".join(cells) + "</tr>")
    return f"""
<div class="qa-wrap">
  <div class="qa-head">
    <div class="qa-title">■ 실적 / 분석</div>
    <div class="qa-unit">(단위 : TON, BT)</div>
  </div>
  <table class="qa-table">
    <thead>{head}</thead>
    <tbody>{"".join(body)}</tbody>
  </table>
</div>
"""


def _qa_build_perf_rows(dash_vals: dict[str, Any]) -> list[dict[str, Any]]:
    """셀 맵 → 실적/분석 표 행."""
    item_rows = [
        ("L-O2", 21, False),
        ("L-N2", 22, False),
        ("L-AR", 23, False),
        ("L-CO2", 24, False),
        ("LPG", 25, False),
    ]
    out: list[dict[str, Any]] = []
    bulk_curr = [0, 0, 0]
    bulk_prior = [0, 0, 0]
    for lab, row, _ in item_rows:
        # D/E/F = m1, G/H/I = m2, J/K/L = m3
        curr = [
            int(_qa_num_or(dash_vals.get(f"D{row}"), 0)),
            int(_qa_num_or(dash_vals.get(f"G{row}"), 0)),
            int(_qa_num_or(dash_vals.get(f"J{row}"), 0)),
        ]
        prior = [
            int(_qa_num_or(dash_vals.get(f"E{row}"), 0)),
            int(_qa_num_or(dash_vals.get(f"H{row}"), 0)),
            int(_qa_num_or(dash_vals.get(f"K{row}"), 0)),
        ]
        rates = [
            str(dash_vals.get(f"F{row}") or "-"),
            str(dash_vals.get(f"I{row}") or "-"),
            str(dash_vals.get(f"L{row}") or "-"),
        ]
        for i in range(3):
            bulk_curr[i] += curr[i]
            bulk_prior[i] += prior[i]
        out.append({"label": lab, "curr": curr, "prior": prior, "rates": rates})
    bulk_rates = [
        _qa_arrow_pct(bulk_curr[i], bulk_prior[i]) for i in range(3)
    ]
    out.append(
        {
            "label": "Bulk 합계",
            "curr": bulk_curr,
            "prior": bulk_prior,
            "rates": bulk_rates,
            "bulk": True,
        }
    )
    out.append(
        {
            "label": "Gas Cylinder",
            "curr": [
                int(_qa_num_or(dash_vals.get("D27"), 0)),
                int(_qa_num_or(dash_vals.get("G27"), 0)),
                int(_qa_num_or(dash_vals.get("J27"), 0)),
            ],
            "prior": [
                int(_qa_num_or(dash_vals.get("E27"), 0)),
                int(_qa_num_or(dash_vals.get("H27"), 0)),
                int(_qa_num_or(dash_vals.get("K27"), 0)),
            ],
            "rates": [
                str(dash_vals.get("F27") or "-"),
                str(dash_vals.get("I27") or "-"),
                str(dash_vals.get("L27") or "-"),
            ],
        }
    )
    return out


def render_quarterly_analysis_tab(
    df: pd.DataFrame | None = None, latest_update_str: str = ""
) -> None:
    """분기 사업분석 — 목표/달성·실적/분석 표 + 분기 버튼 + 목표 입력."""
    st.markdown(
        "<div class='sub-header dashboard-tab-panel-head'>📊 분기 사업분석</div>",
        unsafe_allow_html=True,
    )
    st.markdown(_qa_css(), unsafe_allow_html=True)
    if latest_update_str:
        st.caption(f"대시보드 기준 시각: {latest_update_str}")

    path = ""
    excel_vals: dict[str, Any] = {}
    if load_workbook is not None:
        try:
            path = _qa_ensure_workfile()
            excel_vals = _qa_read_cells(path, [f["cell"] for f in _QA_FIELDS])
        except Exception as err:
            st.warning(f"엑셀 작업파일: {err}")

    today = date.today()
    default_lab = _qa_default_quarter_label()
    dy, dq = _qa_parse_quarter_label(default_lab)
    if "qa_year" not in st.session_state:
        st.session_state["qa_year"] = dy
    if "qa_q" not in st.session_state:
        st.session_state["qa_q"] = dq

    ycol, q1, q2, q3, q4, sp = st.columns([1.1, 0.85, 0.85, 0.85, 0.85, 2.2])
    with ycol:
        year = int(
            st.number_input(
                "연도",
                min_value=2020,
                max_value=2100,
                step=1,
                key="qa_year",
            )
        )
    quarter = int(st.session_state.get("qa_q") or 1)
    for i, col in enumerate((q1, q2, q3, q4), start=1):
        with col:
            st.write("")  # align with number_input
            clicked = st.button(
                f"{i}분기",
                key=f"qa_qbtn_{i}",
                type="primary" if quarter == i else "secondary",
                width="stretch",
            )
            if clicked and quarter != i:
                st.session_state["qa_q"] = i
                st.rerun()
    with sp:
        st.write("")
        st.caption("분기를 누르면 월 집계·표가 바뀝니다.")

    quarter = int(st.session_state.get("qa_q") or 1)
    months = _QA_QUARTER_MONTHS.get(quarter, (1, 2, 3))

    seed_key = f"qa_goal_seed_{year}_{quarter}"
    if st.session_state.get("qa_goal_seed_for") != seed_key:
        st.session_state["qa_goal_seed_for"] = seed_key
        st.session_state["qa_goal_m1"] = int(_qa_num_or(excel_vals.get("E6"), 0))
        st.session_state["qa_goal_m2"] = int(_qa_num_or(excel_vals.get("G6"), 0))
        st.session_state["qa_goal_m3"] = int(_qa_num_or(excel_vals.get("I6"), 0))

    st.markdown("**목표 입력** (백만원 · 표의 「목표」행)")
    # 구분 열 폭에 맞춰 월 입력
    _pad, g1, g2, g3, gsum_c = st.columns([1.35, 1, 1, 1, 1])
    with g1:
        goal_m1 = st.number_input(
            f"{months[0]}월 목표",
            min_value=0,
            max_value=1_000_000,
            step=1,
            key="qa_goal_m1",
        )
    with g2:
        goal_m2 = st.number_input(
            f"{months[1]}월 목표",
            min_value=0,
            max_value=1_000_000,
            step=1,
            key="qa_goal_m2",
        )
    with g3:
        goal_m3 = st.number_input(
            f"{months[2]}월 목표",
            min_value=0,
            max_value=1_000_000,
            step=1,
            key="qa_goal_m3",
        )
    with gsum_c:
        st.metric("목표 합계", f"{int(goal_m1) + int(goal_m2) + int(goal_m3):,}")

    user_goals = {"E6": int(goal_m1), "G6": int(goal_m2), "I6": int(goal_m3)}
    dash_vals = compute_dashboard_values(
        df, year=int(year), months=months, goals=user_goals
    )

    goals_t = (int(goal_m1), int(goal_m2), int(goal_m3))
    sales_t = (
        int(_qa_num_or(dash_vals.get("E7"), 0)),
        int(_qa_num_or(dash_vals.get("G7"), 0)),
        int(_qa_num_or(dash_vals.get("I7"), 0)),
    )
    prior_t = (
        int(_qa_num_or(dash_vals.get("E9"), 0)),
        int(_qa_num_or(dash_vals.get("G9"), 0)),
        int(_qa_num_or(dash_vals.get("I9"), 0)),
    )
    rates_goal = (
        str(dash_vals.get("E8") or "-"),
        str(dash_vals.get("G8") or "-"),
        str(dash_vals.get("I8") or "-"),
        str(dash_vals.get("K8") or "-"),
    )
    rates_prior = (
        str(dash_vals.get("E10") or "-"),
        str(dash_vals.get("G10") or "-"),
        str(dash_vals.get("I10") or "-"),
        str(dash_vals.get("K10") or "-"),
    )

    st.markdown(
        _qa_goal_table_html(
            year, months, goals_t, sales_t, prior_t, rates_goal, rates_prior
        ),
        unsafe_allow_html=True,
    )
    st.markdown(
        _qa_perf_table_html(year, months, _qa_build_perf_rows(dash_vals)),
        unsafe_allow_html=True,
    )

    c1, c2, c3 = st.columns([1.2, 1.2, 2])
    with c1:
        apply = st.button("엑셀에 반영", type="primary", key="qa_apply", width="stretch")
    with c2:
        reset = st.button("템플릿으로 초기화", key="qa_reset", width="stretch")
    with c3:
        if path and os.path.isfile(path):
            with open(path, "rb") as fh:
                st.download_button(
                    "작업 엑셀 다운로드",
                    data=fh.read(),
                    file_name=f"{year}년{quarter}분기사업분석.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="qa_dl",
                )

    if reset:
        if os.path.isfile(QA_TEMPLATE):
            shutil.copy2(QA_TEMPLATE, QA_WORK)
            st.session_state.pop("qa_goal_seed_for", None)
            st.success("템플릿으로 작업 파일을 다시 만들었습니다.")
            st.rerun()
        else:
            st.error("템플릿 파일이 없습니다.")

    if apply:
        if not path:
            st.error("엑셀 작업파일이 없어 반영할 수 없습니다.")
        else:
            cells = [f["cell"] for f in _QA_FIELDS]
            to_write = {
                addr: dash_vals[addr]
                for addr in cells
                if addr in dash_vals and dash_vals[addr] is not None
            }
            try:
                n = _qa_write_cells(path, to_write)
                st.success(f"엑셀에 {n}개 칸을 반영했습니다. (수식·레이아웃 유지)")
                st.rerun()
            except Exception as err:
                st.error(f"반영 실패: {err}")

    st.caption(
        "실적·전년·수량은 대시보드 매출로 집계합니다. "
        "L-O2~LPG 외 품목은 Gas Cylinder로 넣습니다. 합계 수식 칸은 엑셀에서 유지됩니다."
    )
