"""업무일지 비고란(Y~AB) — 원본 한 줄 글자수·읽기/쓰기."""
from __future__ import annotations

import unittest
from datetime import date
from unittest import mock


class WorklogRemarkColumnTest(unittest.TestCase):
    def setUp(self):
        import worklog_tab as wt

        self.wt = wt

    def test_remark_units_match_original_y_ab_width(self):
        self.assertEqual(self.wt._remark_line_units(), 16)
        self.assertEqual(self.wt._hangul_line_limit(self.wt._remark_line_units()), 8)
        self.assertEqual(self.wt._hangul_line_limit(self.wt._client_line_units()), 8)
        self.assertEqual(self.wt._hangul_line_limit(self.wt._content_line_units()), 37)
        label = self.wt._wl_col_limit_label("비고", self.wt._remark_line_units())
        self.assertIn("비고", label)
        self.assertIn("원본 한글 8자", label)

    def test_empty_cells_include_y(self):
        cells = self.wt._empty_cells(date(2026, 9, 12))
        self.assertIn("Y8", cells)
        self.assertIn("Y39", cells)
        self.assertEqual(cells["Y8"], "")

    def test_ensure_body_merges_fills_missing_remark_rows(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.merge_cells("Y8:AB8")
        self.wt._ensure_worklog_body_merges(ws)
        merged = {str(mr) for mr in ws.merged_cells.ranges}
        self.assertIn("Y8:AB8", merged)
        self.assertIn("Y12:AB12", merged)
        self.assertIn("Y39:AB39", merged)
        self.assertIn("C8:F8", merged)
        self.assertIn("G8:X8", merged)
        wb.close()

    def test_pack_and_group_roundtrip_remark(self):
        d = date(2026, 9, 12)
        entries = [{
            "client": "거래처A",
            "content": "점검",
            "remarks": "방문",
            "blank_after": 1,
        }]
        cells = self.wt._pack_entries_to_cells(d, entries)
        self.assertEqual(cells["C8"], "거래처A")
        self.assertEqual(cells["G8"], "점검")
        self.assertEqual(cells["Y8"], "방문")
        grouped = self.wt._grouped_entries_from_cells(cells)
        self.assertEqual(len(grouped), 1)
        self.assertEqual(grouped[0]["remarks"].strip(), "방문")
        self.assertEqual(grouped[0]["remark_lines"][0], "방문")

    def test_remark_only_row_is_kept(self):
        d = date(2026, 9, 12)
        cells = self.wt._empty_cells(d)
        cells["Y8"] = "전화"
        grouped = self.wt._grouped_entries_from_cells(cells)
        self.assertEqual(len(grouped), 1)
        self.assertEqual(grouped[0]["remarks"], "전화")
        self.assertTrue(self.wt._worklog_cells_have_draft(cells))

    def test_remark_overflow_splits_to_next_line(self):
        too_long = "가" * 12  # 24 units > 16
        chunks = self.wt._chunk_text(too_long, self.wt._remark_line_units())
        self.assertGreaterEqual(len(chunks), 2)
        self.assertLessEqual(self.wt._display_units(chunks[0]), self.wt._remark_line_units())

    def test_spaces_follow_original_cell_width(self):
        self.assertAlmostEqual(self.wt._char_units(" "), 0.58)
        self.assertEqual(self.wt._char_units("가"), 2)
        self.assertEqual(self.wt._orig_cell_px("content"), 666)
        self.assertAlmostEqual(self.wt._look_font_pt(), 14.0 * self.wt._WL_INPUT_LOOK_SCALE)
        self.assertGreater(self.wt._WL_INPUT_LOOK_SCALE, self.wt._WL_PREVIEW_SCALE)
        self.assertEqual(self.wt._input_cell_px("content"), int(round(666 * self.wt._WL_INPUT_LOOK_SCALE)))
        self.assertEqual(self.wt._orig_cell_px("client"), 148)
        self.assertEqual(self.wt._orig_cell_px("remark"), self.wt._orig_cell_px("client"))
        self.assertEqual(self.wt._input_cell_px("client"), self.wt._input_cell_px("remark"))
        self.assertEqual(
            self.wt._input_cell_px("client"),
            int(round(148 * self.wt._WL_INPUT_LOOK_SCALE))
            + int(self.wt._WL_INPUT_SIDE_CHROME_PX)
            + int(self.wt._WL_INPUT_SIDE_WIDEN_PX),
        )
        eight = "가" * 8
        nine = "가" * 9
        self.assertEqual(len(self.wt._chunk_text(eight, self.wt._remark_line_units())), 1)
        self.assertGreaterEqual(len(self.wt._chunk_text(nine, self.wt._remark_line_units())), 2)
        max_u = self.wt._content_line_units()
        hangul_ok = "가" * 37
        hangul_over = "가" * 38
        self.assertLessEqual(self.wt._display_units(hangul_ok), max_u)
        self.assertGreater(self.wt._display_units(hangul_over), max_u)
        self.assertEqual(len(self.wt._chunk_text(hangul_ok, max_u)), 1)
        self.assertGreaterEqual(len(self.wt._chunk_text(hangul_over, max_u)), 2)
        # 띄어쓰기는 한글보다 좁아 원본 칸 끝까지 더 쓴 뒤에 다음 칸으로 간다.
        pairs_fit = "가 " * 28
        pairs_over = "가 " * 29
        self.assertLessEqual(self.wt._display_units(pairs_fit), max_u)
        self.assertGreater(self.wt._display_units(pairs_over), max_u)
        self.assertEqual(len(self.wt._chunk_text(pairs_fit, max_u)), 1)
        self.assertGreaterEqual(len(self.wt._chunk_text(pairs_over, max_u)), 2)

    def test_mid_start_counts_space_like_preview(self):
        max_u = self.wt._content_line_units()
        indent = " " * 12
        over = indent + ("가" * 37)
        self.assertGreater(self.wt._display_units(over), self.wt._display_units("가" * 37))
        self.assertGreater(self.wt._display_units(over), max_u)
        head, tail = self.wt._fit_by_units(over, max_u)
        self.assertLessEqual(self.wt._display_units(head), max_u)
        self.assertTrue(tail.startswith("가"))
        self.assertFalse(tail.startswith(" "))
        self.assertGreaterEqual(len(self.wt._chunk_text(over, max_u)), 2)

    def test_ui_shows_original_limits(self):
        with open(self.wt.__file__, encoding="utf-8") as f:
            src = f.read()
        self.assertIn('_wl_col_limit_label("거래처"', src)
        self.assertIn('_wl_col_limit_label("내용"', src)
        self.assertIn('_wl_col_limit_label("비고"', src)
        self.assertIn('"variant": "remark"', src)
        self.assertNotIn("제미나이", src)
        self.assertNotIn("worklog_google_ai_search", src)
        self.assertNotIn("_mount_google_ai_search", src)
        self.assertIn('st.text_area("익일업무"', src)
        self.assertIn('st.text_area("특이사항"', src)
        self.assertIn("fixed_rows", src)
        self.assertNotIn("＋ 항목 추가", src)


class WorklogSheetBodyTest(unittest.TestCase):
    def setUp(self):
        import worklog_tab as wt

        self.wt = wt

    def test_pad_sheet_is_32_and_excludes_next_notes_rows(self):
        self.assertEqual(self.wt.WL_SHEET_N, 32)
        self.assertEqual(self.wt.WL_CONTENT_ROWS[-1], 39)
        self.assertEqual(self.wt.WL_NEXT_ROWS, [40, 41, 42, 43])
        self.assertEqual(self.wt.WL_NOTE_ROWS, [44, 45, 46, 47])
        padded = self.wt._pad_sheet_lines(["A", "B"])
        self.assertEqual(len(padded), 32)
        self.assertEqual(padded[0], "A")
        self.assertEqual(padded[1], "B")
        self.assertEqual(padded[2], "")

    def test_two_clients_without_blank_are_two_summary_items(self):
        d = date(2026, 9, 15)
        clients = [""] * 32
        contents = [""] * 32
        clients[0], contents[0] = "거래처갑", "점검"
        clients[1], contents[1] = "거래처을", "보수"
        cells = self.wt._pack_sheet_to_cells(d, clients, contents, [""] * 32)
        grouped = self.wt._grouped_entries_from_cells(cells)
        self.assertEqual(len(grouped), 2)
        rows = [self.wt._summary_row_from_entry(e) for e in grouped]
        self.assertEqual(rows[0][0], "거래처갑")
        self.assertEqual(rows[1][0], "거래처을")

    def test_next_and_notes_stay_in_d40_d47_not_body(self):
        d = date(2026, 9, 15)
        cells = self.wt._pack_sheet_to_cells(
            d,
            ["본문거래처"] + [""] * 31,
            ["본문내용"] + [""] * 31,
            [""] * 32,
            next_day=["익일할일"],
            notes=["특이메모"],
        )
        self.assertEqual(cells["C8"], "본문거래처")
        self.assertEqual(cells["G8"], "본문내용")
        self.assertEqual(cells["C39"], "")
        self.assertEqual(cells["G39"], "")
        self.assertIn("익일할일", cells["D40"])
        self.assertIn("특이메모", cells["D44"])
        for r in self.wt.WL_CONTENT_ROWS:
            self.assertNotIn("익일할일", str(cells.get(f"G{r}", "")))
            self.assertNotIn("특이메모", str(cells.get(f"G{r}", "")))
        _, next_day, notes = self.wt._entries_from_cells(cells)
        self.assertEqual(next_day[0], "익일할일")
        self.assertEqual(notes[0], "특이메모")

    def test_trailing_empty_sheet_rows_do_not_count_as_used(self):
        usage = self.wt._sheet_row_usage(
            ["갑", ""] + [""] * 30,
            ["내용", ""] + [""] * 30,
            [""] * 32,
        )
        self.assertEqual(usage["used"], 1)
        self.assertEqual(usage["remaining"], 31)
        self.assertFalse(usage["overflow"])


class WorklogInputPagesTest(unittest.TestCase):
    def setUp(self):
        import worklog_tab as wt

        self.wt = wt

    def test_extra_sheet_names_are_not_calendar_days(self):
        d = date(2026, 9, 16)
        self.assertEqual(self.wt._extra_page_sheet_name(2), "p2")
        self.assertEqual(self.wt._archive_extra_sheet_name(d, 2), "16p2")
        self.assertEqual(self.wt._extra_page_n_from_sheet_name("p3"), 3)
        self.assertEqual(self.wt._extra_page_n_from_sheet_name("16p2"), 2)
        self.assertIsNone(self.wt._extra_page_n_from_sheet_name("16"))
        self.assertIsNone(self.wt._extra_page_n_from_sheet_name("2026-09-16"))
        self.assertTrue(self.wt._is_day_extra_sheet_name("p2"))
        self.assertTrue(self.wt._is_archive_extra_sheet_name("16p2", d))
        self.assertFalse(self.wt._is_archive_extra_sheet_name("16p2", date(2026, 9, 1)))

    def test_summary_includes_extra_pages_same_list(self):
        d = date(2026, 9, 16)
        page1 = self.wt._pack_sheet_to_cells(
            d, ["갑"] + [""] * 31, ["1페이지내용"] + [""] * 31, [""] * 32, ["익일"], ["특이"],
        )
        page2 = self.wt._pack_sheet_to_cells(
            d, ["을"] + [""] * 31, ["2페이지내용"] + [""] * 31, [""] * 32, [], [],
        )
        html = self.wt.render_readable_preview_html(d, self.wt._attach_extra_pages(page1, [page2]))
        self.assertIn("갑", html)
        self.assertIn("1페이지내용", html)
        self.assertIn("을", html)
        self.assertIn("2페이지내용", html)
        self.assertIn("익일", html)
        self.assertIn("특이", html)
        self.assertEqual(html.count("<h3>익일업무</h3>"), 1)
        self.assertEqual(html.count("<h3>특 이 사 항</h3>"), 1)

    def test_ui_keeps_sheet_layout_and_adds_page_bar(self):
        with open(self.wt.__file__, encoding="utf-8") as f:
            src = f.read()
        self.assertIn('f"{_pi + 1}페이지"', src)
        self.assertIn('on_click=_queue_worklog_add', src)
        self.assertIn("_mount_entry_client_editor(iso2, pi, _cu)", src)
        self.assertIn("_mount_entry_lines_editor(iso2, pi, max_u)", src)
        self.assertIn("_mount_entry_remark_editor(iso2, pi, _ru)", src)
        self.assertIn("st.columns([_side, _gw, _side]", src)
        self.assertIn('wl_page_bar_{iso2}', src)
        self.assertIn("horizontal=True", src)
        self.assertIn("2.4rem", src)
        self.assertIn('st.text_area("익일업무"', src)
        self.assertIn('st.text_area("특이사항"', src)
        self.assertNotIn("＋ 항목 추가", src)

    def test_hidden_page_keeps_snapshot_when_live_is_wiped(self):
        from unittest.mock import patch

        iso = "2026-09-16"
        ss = {}
        with patch.object(self.wt.st, "session_state", ss):
            ss[self.wt._page_count_key(iso)] = 2
            ss[self.wt._page_idx_key(iso)] = 1
            ss[self.wt._page_snap_key(iso, 0)] = (
                ["갑"] + [""] * 31,
                ["본문1"] + [""] * 31,
                [""] * 32,
            )
            ss[self.wt._entry_clients_live_key(iso, 0)] = [""] * 32
            ss[self.wt._entry_lines_live_key(iso, 0)] = [""] * 32
            ss[self.wt._entry_remarks_live_key(iso, 0)] = [""] * 32
            clients, contents, remarks = self.wt._sheet_lines_from_widgets_at(iso, 0)
            self.assertEqual(clients[0], "갑")
            self.assertEqual(contents[0], "본문1")
            self.assertEqual(remarks[0], "")

    def test_write_read_extra_page_roundtrip(self):
        import os
        import tempfile

        d = date(2026, 9, 16)
        page1 = self.wt._pack_sheet_to_cells(
            d, ["갑"] + [""] * 31, ["본문1"] + [""] * 31, [""] * 32, ["익일1"], [],
        )
        page2 = self.wt._pack_sheet_to_cells(
            d, ["을"] + [""] * 31, ["본문2"] + [""] * 31, [""] * 32, [], [],
        )
        cells = self.wt._attach_extra_pages(page1, [page2])
        with tempfile.TemporaryDirectory() as tmp:
            cache = os.path.join(tmp, "worklog")
            os.makedirs(cache, exist_ok=True)
            tpl = os.path.join(cache, "template.xlsx")
            from openpyxl import Workbook, load_workbook

            wb = Workbook()
            wb.active["C5"] = ""
            wb.save(tpl)
            wb.close()
            with mock.patch.object(self.wt, "WORKLOG_DIR", cache), mock.patch.object(self.wt, "WORKLOG_TEMPLATE", tpl):
                path = os.path.join(cache, f"{d.isoformat()}.xlsx")
                self.wt.write_cells_to_path(path, d, cells, force_template=True)
                wb2 = load_workbook(path)
                try:
                    self.assertIn("p2", wb2.sheetnames)
                    self.assertEqual(str(wb2.active["G8"].value or ""), "본문1")
                    self.assertEqual(str(wb2["p2"]["G8"].value or ""), "본문2")
                    self.assertEqual(str(wb2["p2"]["C8"].value or ""), "을")
                    self.assertFalse(str(wb2["p2"]["D40"].value or "").strip())
                finally:
                    wb2.close()
                wb3 = load_workbook(path)
                try:
                    extras = self.wt._extra_page_cells_from_workbook(wb3, d)
                    self.assertEqual(len(extras), 1)
                    self.assertEqual(extras[0]["G8"], "본문2")
                finally:
                    wb3.close()


if __name__ == "__main__":
    unittest.main()
